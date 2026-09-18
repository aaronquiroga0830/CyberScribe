"""Regression coverage for DOCX ingestion without optional unstructured packages."""
import tempfile
import unittest
from pathlib import Path

from docx import Document

from src.ingest.loaders import load_file, load_mission_documents


class TestDocumentLoaders(unittest.TestCase):
    def test_docx_preserves_paragraph_table_order_and_source(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "findings.docx"
            document = Document()
            document.add_paragraph("Before: café")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Finding"
            table.cell(0, 1).text = "Risk"
            table.cell(1, 0).text = "Exposed RDP"
            table.cell(1, 1).text = "High"
            document.add_paragraph("After: restrict access")
            document.save(path)

            loaded = load_file(path)
            self.assertEqual(len(loaded), 1)
            self.assertEqual(
                loaded[0].page_content,
                "Before: café\nFinding\tRisk\nExposed RDP\tHigh\nAfter: restrict access",
            )
            self.assertEqual(loaded[0].metadata["source"], str(path))
            self.assertEqual(loaded[0].metadata["mission_file"], path.name)

    def test_nested_mission_sources_receive_document_type(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            logs = root / "crew_logs"
            logs.mkdir()
            (logs / "day1.txt").write_text("09:00 - Checked gateway", encoding="utf-8")
            document = Document()
            document.add_paragraph("RDP is exposed")
            document.save(root / "findings.docx")

            loaded = load_mission_documents(root)
            self.assertEqual(len(loaded), 2)
            self.assertEqual({d.metadata["doc_type"] for d in loaded}, {"crew_log", "findings"})

    def test_missing_or_empty_source_folder(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.assertEqual(load_mission_documents(root), [])
            self.assertEqual(load_mission_documents(root / "missing"), [])
            self.assertEqual(load_file(root / "missing.docx"), [])
