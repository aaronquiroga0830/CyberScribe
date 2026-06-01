"""
Tests for report .docx output and HTML conversion.
Run from project root: python -m unittest tests.test_report_docx -v
"""
import tempfile
import unittest
from pathlib import Path


class TestHtmlToDocx(unittest.TestCase):
    def test_plain_text(self):
        """Plain text (no HTML) is written as a single paragraph."""
        from src.html_to_docx import html_to_docx
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.docx"
            html_to_docx("Hello world.\nSecond line.", out)
            self.assertTrue(out.exists())
            self.assertGreater(out.stat().st_size, 0)

    def test_empty(self):
        """Empty content produces a minimal .docx."""
        from src.html_to_docx import html_to_docx
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.docx"
            html_to_docx("", out)
            self.assertTrue(out.exists())
            self.assertGreater(out.stat().st_size, 0)

    def test_with_margins(self):
        """Margins are applied to the section."""
        from src.html_to_docx import html_to_docx
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.docx"
            html_to_docx("<p>Test</p>", out, margins={"top": 1.5, "right": 1, "bottom": 1, "left": 1})
            self.assertTrue(out.exists())
            from docx import Document
            doc = Document(str(out))
            self.assertAlmostEqual(doc.sections[0].top_margin.inches, 1.5, places=2)

    def test_simple_html(self):
        """Simple HTML is converted to .docx."""
        from src.html_to_docx import html_to_docx
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.docx"
            html_to_docx("<p>Para one</p><p>Para two</p>", out)
            self.assertTrue(out.exists())
            from docx import Document
            doc = Document(str(out))
            paras = list(doc.paragraphs)
            self.assertGreaterEqual(len(paras), 2)
            self.assertIn("Para one", paras[0].text)
            self.assertIn("Para two", paras[1].text)


if __name__ == "__main__":
    unittest.main()
